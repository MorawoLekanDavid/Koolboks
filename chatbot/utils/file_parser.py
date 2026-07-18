"""Extract plain text from uploaded knowledge base documents."""
import io


def parse_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace").strip()


def parse_md(data: bytes) -> str:
    return data.decode("utf-8", errors="replace").strip()


def parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(p.strip() for p in pages if p.strip())
    except Exception as e:
        raise ValueError(f"Could not parse PDF: {e}")


def parse_docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Could not parse DOCX: {e}")


def parse_xlsx(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"### {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c for c in cells):
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    except Exception as e:
        raise ValueError(f"Could not parse XLSX: {e}")


PARSERS = {
    "txt":  parse_txt,
    "md":   parse_md,
    "pdf":  parse_pdf,
    "docx": parse_docx,
    "xlsx": parse_xlsx,
}


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Return (file_type, extracted_text). Raises ValueError for unsupported types."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in PARSERS:
        raise ValueError(f"Unsupported file type '.{ext}'. Allowed: {', '.join(PARSERS)}")
    return ext, PARSERS[ext](data)
